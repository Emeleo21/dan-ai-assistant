import io

import streamlit as st
from docx import Document


def create_word_document(title: str, content: str, filename: str = "document.docx") -> str:
    """Create an editable Word (.docx) document with a title and body text.
    Separate paragraphs with blank lines. Lines starting with '# ' become headings.
    Use this when the user asks for a Word file, .docx, or an editable document.
    The user will get a download button automatically."""
    try:
        if not filename.lower().endswith(".docx"):
            filename += ".docx"

        doc = Document()
        doc.add_heading(title, level=0)

        for block in content.split("\n\n"):
            block = block.strip()
            if not block:
                continue
            if block.startswith("# "):
                doc.add_heading(block[2:], level=1)
            else:
                doc.add_paragraph(block)

        buffer = io.BytesIO()
        doc.save(buffer)

        st.session_state.setdefault("pending_files", []).append({
            "name": filename,
            "data": buffer.getvalue(),
            "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        })
        return f"Word document '{filename}' created. Tell the user it is ready to download below."
    except Exception as e:
        return f"Error creating Word document: {e}"